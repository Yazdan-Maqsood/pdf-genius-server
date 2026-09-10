#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
PDF to PowerPoint Converter
Two modes:
  - text:  Editable text, no theme (smaller file)
  - image: Preserved theme/design as image (larger file)

Usage: python pdf_to_ppt.py <input.pdf> <output.pptx> [mode]
Mode: text (default) | image
"""

import sys
import os
import json
import traceback
import platform
import io


def is_scanned_pdf(pdf_path):
    """Check if PDF is scanned (no extractable text)."""
    try:
        import pymupdf
        doc = pymupdf.open(pdf_path)
        
        total_text = ''
        for i in range(min(3, len(doc))):
            page = doc[i]
            text = page.get_text().strip()
            total_text += text
        
        doc.close()
        return len(total_text) < 50
    except Exception:
        return False


def render_pdf_pages_as_images(pdf_path, output_dir, dpi=150):
    """Render each PDF page as a high-quality PNG image."""
    try:
        import pymupdf
        
        doc = pymupdf.open(pdf_path)
        image_paths = []
        
        scale = dpi / 72.0
        
        for page_num in range(len(doc)):
            page = doc[page_num]
            mat = pymupdf.Matrix(scale, scale)
            pix = page.get_pixmap(matrix=mat, alpha=False)
            
            image_path = os.path.join(output_dir, f"page_{page_num + 1}.png")
            pix.save(image_path)
            image_paths.append(image_path)
            
            print(json.dumps({"info": f"Rendered page {page_num + 1}/{len(doc)}"}), file=sys.stderr)
        
        doc.close()
        return image_paths
        
    except Exception as e:
        print(json.dumps({"error": f"Render failed: {str(e)}"}), file=sys.stderr)
        return []


def extract_text_with_ocr(pdf_path):
    """Extract text from scanned PDF using OCR."""
    try:
        import pymupdf
        import pytesseract
        from PIL import Image
        
        if platform.system() == 'Windows':
            possible_paths = [
                r'C:\Program Files\Tesseract-OCR\tesseract.exe',
                r'C:\Program Files (x86)\Tesseract-OCR\tesseract.exe',
            ]
            for path in possible_paths:
                if os.path.exists(path):
                    pytesseract.pytesseract.tesseract_cmd = path
                    break
        
        doc = pymupdf.open(pdf_path)
        pages_text = []
        
        for page_num in range(len(doc)):
            page = doc[page_num]
            pix = page.get_pixmap(matrix=pymupdf.Matrix(2, 2))
            img_data = pix.tobytes("png")
            img = Image.open(io.BytesIO(img_data))
            
            text = pytesseract.image_to_string(img)
            pages_text.append(text)
            
            print(json.dumps({"info": f"OCR page {page_num + 1}/{len(doc)}"}), file=sys.stderr)
        
        doc.close()
        return pages_text
        
    except Exception as e:
        print(json.dumps({"error": f"OCR failed: {str(e)}"}), file=sys.stderr)
        return []


def extract_text_with_pdfplumber(pdf_path):
    """Extract text from text-based PDF."""
    try:
        import pdfplumber
        
        pages_text = []
        with pdfplumber.open(pdf_path) as pdf:
            for page_num, page in enumerate(pdf.pages, 1):
                text = page.extract_text() or ''
                pages_text.append(text)
                print(json.dumps({"info": f"Extracted page {page_num}"}), file=sys.stderr)
        
        return pages_text
    except Exception as e:
        print(json.dumps({"error": f"pdfplumber failed: {str(e)}"}), file=sys.stderr)
        return []


def create_text_mode_ppt(pdf_path, output_path):
    """Create PPT with editable text (no theme/design)."""
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    
    # Detect scanned PDF
    scanned = is_scanned_pdf(pdf_path)
    
    # Extract text
    if scanned:
        pages_text = extract_text_with_ocr(pdf_path)
    else:
        pages_text = extract_text_with_pdfplumber(pdf_path)
    
    if not pages_text:
        return {"success": False, "error": "Could not extract text from PDF"}
    
    # Create presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    slide_count = 0
    
    for page_num, page_text in enumerate(pages_text, 1):
        if not page_text or not page_text.strip():
            continue
        
        blank_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_layout)
        slide_count += 1
        
        # Extract title
        lines = [line.strip() for line in page_text.split('\n') if line.strip()]
        title_text = f"Page {page_num}"
        content_lines = lines
        
        if lines:
            first_line = lines[0]
            if len(first_line) < 80:
                title_text = first_line
                content_lines = lines[1:]
        
        # Add title
        title_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(0.3), Inches(12.33), Inches(1)
        )
        title_frame = title_box.text_frame
        title_frame.word_wrap = True
        title_para = title_frame.paragraphs[0]
        title_para.text = title_text[:200]
        title_para.font.size = Pt(28)
        title_para.font.bold = True
        title_para.font.color.rgb = RGBColor(0x66, 0x7E, 0xEA)
        
        # Add content
        if content_lines:
            content_box = slide.shapes.add_textbox(
                Inches(0.5), Inches(1.5), Inches(12.33), Inches(5.5)
            )
            content_frame = content_box.text_frame
            content_frame.word_wrap = True
            
            # Group into paragraphs
            paragraphs = []
            current = []
            
            for line in content_lines:
                if len(line) < 100 and (line.endswith(':') or line.endswith('.')):
                    if current:
                        paragraphs.append(' '.join(current))
                        current = []
                    current.append(line)
                else:
                    current.append(line)
                
                if len(' '.join(current)) > 300:
                    paragraphs.append(' '.join(current))
                    current = []
            
            if current:
                paragraphs.append(' '.join(current))
            
            for idx, para in enumerate(paragraphs[:15]):
                if idx == 0:
                    p = content_frame.paragraphs[0]
                else:
                    p = content_frame.add_paragraph()
                
                p.text = para[:500]
                p.font.size = Pt(14)
                p.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
                p.space_after = Pt(8)
        
        # Add slide number
        footer_box = slide.shapes.add_textbox(
            Inches(11.5), Inches(7), Inches(1.5), Inches(0.4)
        )
        footer_frame = footer_box.text_frame
        footer_para = footer_frame.paragraphs[0]
        footer_para.text = f"{slide_count} / {len(pages_text)}"
        footer_para.font.size = Pt(10)
        footer_para.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
        footer_para.alignment = PP_ALIGN.RIGHT
    
    if slide_count == 0:
        blank_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_layout)
        textbox = slide.shapes.add_textbox(Inches(2), Inches(3), Inches(9), Inches(1))
        p = textbox.text_frame.paragraphs[0]
        p.text = "No content could be extracted"
        p.font.size = Pt(20)
    
    prs.save(output_path)
    
    return {
        "success": True,
        "slides": slide_count,
        "method": "text" + ("_ocr" if scanned else "")
    }


def create_image_mode_ppt(pdf_path, output_path, output_dir):
    """Create PPT with preserved design (image backgrounds)."""
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    import pymupdf
    
    # Create temp dir for images
    temp_img_dir = os.path.join(output_dir, f"temp_img_{os.getpid()}")
    os.makedirs(temp_img_dir, exist_ok=True)
    
    # Render pages as images
    image_paths = render_pdf_pages_as_images(pdf_path, temp_img_dir, dpi=150)
    
    if not image_paths:
        return {"success": False, "error": "Could not render PDF pages"}
    
    # Get PDF page dimensions
    doc = pymupdf.open(pdf_path)
    page_dimensions = []
    for page in doc:
        rect = page.rect
        page_dimensions.append((rect.width, rect.height))
    doc.close()
    
    # Create presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    slide_count = 0
    
    for idx, img_path in enumerate(image_paths):
        if idx >= len(page_dimensions):
            break
        
        page_width, page_height = page_dimensions[idx]
        page_ratio = page_width / page_height
        
        slide_width = prs.slide_width
        slide_height = prs.slide_height
        slide_ratio = slide_width / slide_height
        
        # Fit image to slide
        if page_ratio > slide_ratio:
            img_width = slide_width
            img_height = int(slide_width / page_ratio)
            img_left = 0
            img_top = int((slide_height - img_height) / 2)
        else:
            img_height = slide_height
            img_width = int(slide_height * page_ratio)
            img_left = int((slide_width - img_width) / 2)
            img_top = 0
        
        # Create blank slide
        blank_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_layout)
        slide_count += 1
        
        # ✅ Add page image as slide background
        slide.shapes.add_picture(
            img_path, img_left, img_top,
            width=img_width, height=img_height
        )
        
        # Add slide number
        footer_box = slide.shapes.add_textbox(
            Inches(11.5), Inches(7), Inches(1.5), Inches(0.4)
        )
        footer_frame = footer_box.text_frame
        footer_para = footer_frame.paragraphs[0]
        footer_para.text = f"{slide_count} / {len(image_paths)}"
        footer_para.font.size = Pt(10)
        footer_para.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        footer_para.alignment = PP_ALIGN.RIGHT
    
    prs.save(output_path)
    
    # Cleanup
    try:
        import shutil
        shutil.rmtree(temp_img_dir, ignore_errors=True)
    except Exception:
        pass
    
    return {
        "success": True,
        "slides": slide_count,
        "method": "image"
    }


def convert_pdf_to_ppt(input_path, output_path, mode="text"):
    """Main conversion function."""
    try:
        # Validate input
        if not os.path.exists(input_path):
            return {"success": False, "error": f"Input not found: {input_path}"}
        
        # Validate mode
        if mode not in ["text", "image"]:
            mode = "text"
        
        # Ensure python-pptx
        try:
            from pptx import Presentation
        except ImportError:
            return {
                "success": False,
                "error": "python-pptx not installed. Run: pip install python-pptx"
            }
        
        # Create output dir
        output_dir = os.path.dirname(output_path)
        if output_dir and not os.path.exists(output_dir):
            os.makedirs(output_dir, exist_ok=True)
        
        print(json.dumps({"info": f"Mode: {mode}"}), file=sys.stderr)
        
        # Run appropriate conversion
        if mode == "image":
            result = create_image_mode_ppt(input_path, output_path, output_dir)
        else:
            result = create_text_mode_ppt(input_path, output_path)
        
        if not result.get("success"):
            return result
        
        # Verify output
        if not os.path.exists(output_path):
            return {"success": False, "error": "Output file not created"}
        
        size_kb = os.path.getsize(output_path) / 1024
        
        return {
            "success": True,
            "output_path": output_path,
            "size_kb": round(size_kb, 2),
            "slides": result.get("slides", 0),
            "mode": mode,
            "method": result.get("method", mode)
        }
        
    except Exception as e:
        return {
            "success": False,
            "error": str(e),
            "traceback": traceback.format_exc()
        }


def main():
    if len(sys.argv) < 3:
        print(json.dumps({
            "success": False,
            "error": "Usage: python pdf_to_ppt.py <input.pdf> <output.pptx> [mode]"
        }))
        sys.exit(1)
    
    input_path = sys.argv[1]
    output_path = sys.argv[2]
    mode = sys.argv[3] if len(sys.argv) > 3 else "text"
    
    result = convert_pdf_to_ppt(input_path, output_path, mode)
    print(json.dumps(result))
    sys.exit(0 if result.get("success") else 1)


if __name__ == "__main__":
    main()